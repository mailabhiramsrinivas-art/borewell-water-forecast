#!/usr/bin/env python3
"""Build the polished Word handoff from the canonical Markdown source."""

from __future__ import annotations

import re
import json, hashlib, os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/Bengaluru_Borewell_Model_Mathematical_Handoff.md"
OUTPUT = ROOT / "docs/Bengaluru_Borewell_Model_Mathematical_Handoff.docx"

NAVY = "17324D"
BLUE = "1F6F8B"
PALE_BLUE = "EAF3F7"
LIGHT_BLUE = "F4F8FA"
MID_GRAY = "D8E0E5"
TEXT = RGBColor(34, 43, 50)
MUTED = RGBColor(88, 102, 112)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=100, bottom=65, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    props.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        props.append(u)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _balanced_group(text: str, start: int) -> tuple[str, int] | None:
    if start >= len(text) or text[start] != "{":
        return None
    depth = 0
    for index in range(start, len(text)):
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:index], index + 1
    return None


def _replace_fractions(text: str) -> str:
    while "\\frac" in text:
        start = text.find("\\frac")
        cursor = start + len("\\frac")
        while cursor < len(text) and text[cursor].isspace():
            cursor += 1
        numerator = _balanced_group(text, cursor)
        if not numerator:
            text = text[:start] + "/" + text[cursor:]
            continue
        denominator_start = numerator[1]
        while denominator_start < len(text) and text[denominator_start].isspace():
            denominator_start += 1
        denominator = _balanced_group(text, denominator_start)
        if not denominator:
            text = text[:start] + f"({numerator[0]})/" + text[numerator[1]:]
            continue
        replacement = f"({numerator[0]})/({denominator[0]})"
        text = text[:start] + replacement + text[denominator[1]:]
    return text


def linearize_math(text: str) -> str:
    """Convert embedded LaTeX into readable, reconstructable linear math."""
    text = text.replace("$", "")
    text = _replace_fractions(text)
    text = re.sub(r"\\hat\s*\{([^{}]+)\}", r"hat(\1)", text)
    text = re.sub(r"\\hat\s+([A-Za-z]+)", r"hat(\1)", text)
    text = re.sub(r"\\tilde\s*\{([^{}]+)\}", r"tilde(\1)", text)
    text = re.sub(r"\\tilde\s+([A-Za-z]+)", r"tilde(\1)", text)
    text = re.sub(r"\\bar\s*\{([^{}]+)\}", r"bar(\1)", text)
    text = re.sub(r"\\bar\s+([A-Za-z]+)", r"bar(\1)", text)
    text = re.sub(r"\\mathcal\s*\{([^{}]+)\}", r"\1", text)
    text = re.sub(r"\\text\s*\{([^{}]+)\}", r"\1", text)
    replacements = {
        r"\Delta": "Δ", r"\sigma": "σ", r"\beta": "β", r"\lambda": "λ",
        r"\mu": "μ", r"\pi": "π", r"\tau": "τ", r"\epsilon": "ε",
        r"\omega": "ω", r"\le": "≤", r"\ge": "≥", r"\times": "×",
        r"\rightarrow": "→", r"\in": "∈", r"\sum": "Σ", r"\sqrt": "√", r"\ldots": "…",
        r"\quad": "   ", r"\;": " ", r"\,": " ", r"\!": "",
        r"\left": "", r"\right": "", r"\begin{cases}": "{ ", r"\end{cases}": "",
        r"\{": "{", r"\}": "}", r"\\": "  |  ", r"\#": "#", r"\ln": "ln", r"\log": "log",
        r"\sin": "sin", r"\cos": "cos", r"\min": "min", r"\max": "max",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"_\{([^{}]+)\}", r"_(\1)", text)
    text = re.sub(r"\^\{([^{}]+)\}", r"^(\1)", text)
    text = text.replace("&", " ")
    text = re.sub(r"\\\s+", " ", text)
    text = text.replace("{", "(").replace("}", ")")
    text = re.sub(r"\\([A-Za-z]+)", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def add_inline(paragraph, text: str, *, default_bold=False, size=None, color=None) -> None:
    """Add a small useful Markdown subset: bold, inline code, and bare URLs."""
    text = linearize_math(text)
    token_re = re.compile(r"(\*\*.+?\*\*|`[^`]+`|https?://[^\s)]+)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.bold = default_bold
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt((size or 9.4) - 0.3)
            run.font.color.rgb = RGBColor(23, 69, 84)
            run.font.highlight_color = None
        else:
            add_hyperlink(paragraph, token, token)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = default_bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 27, "000000", 0, 10),
        ("Heading 1", 17, "000000", 14, 7),
        ("Heading 2", 13, "000000", 10, 5),
        ("Heading 3", 10.7, "000000", 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9.5)
    styles["List Bullet"].paragraph_format.left_indent = Cm(0.6)
    styles["List Bullet"].paragraph_format.first_line_indent = Cm(-0.28)
    styles["List Bullet"].paragraph_format.space_after = Pt(2.5)

    styles["List Number"].font.name = "Aptos"
    styles["List Number"].font.size = Pt(9.5)
    styles["List Number"].paragraph_format.left_indent = Cm(0.7)
    styles["List Number"].paragraph_format.first_line_indent = Cm(-0.35)
    styles["List Number"].paragraph_format.space_after = Pt(2.5)

    if "Equation Display" not in styles:
        equation = styles.add_style("Equation Display", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = styles["Equation Display"]
    equation.font.name = "Cambria Math"
    equation.font.size = Pt(9.6)
    equation.font.color.rgb = TEXT
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.keep_together = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Aptos Mono"
    code.font.size = Pt(8.2)
    code.font.color.rgb = RGBColor(38, 54, 65)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.space_before = Pt(1)
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.keep_together = False


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.65)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    header = section.header.paragraphs[0]
    header.text = "BENGALURU BOREWELL FORECASTING POC  •  MATHEMATICAL HANDOFF"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.4)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0,0,0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("26 SEPTEMBER 2026   •   ")
    run.font.name = "Aptos"
    run.font.size = Pt(7.4)
    run.font.color.rgb = MUTED
    add_page_number(footer)


def add_cover(doc: Document, headings: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(30)
    run = p.add_run("PROJECT TECHNICAL HANDOFF")
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0,0,0)
    run.font.letter_spacing = Pt(1.2)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(12)
    title.add_run("Mathematical and Technical Handoff for the Bengaluru Borewell Forecasting Proof of Concept")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_inline(subtitle, "Complete data, model, validation, website, rainfall, and deployment specification", default_bold=True, size=12, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    add_inline(p, "Revised 26 September 2026", default_bold=True, size=9.5, color=RGBColor.from_string(BLUE))

    purpose = doc.add_paragraph()
    purpose.paragraph_format.space_before = Pt(16)
    purpose.paragraph_format.space_after = Pt(12)
    purpose.paragraph_format.left_indent = Cm(0.4)
    purpose.paragraph_format.right_indent = Cm(2.5)
    add_inline(
        purpose,
        "A self-contained record of what was supplied, what was calculated, what was validated, what the public website displays, and what remains scientifically unsupported.",
        size=11,
        color=TEXT,
    )

    scope = doc.add_paragraph()
    scope.paragraph_format.space_before = Pt(26)
    scope.add_run("SCOPE IN ONE SENTENCE\n").bold = True
    scope.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
    scope.runs[0].font.size = Pt(8)
    run = scope.add_run("Real session transitions from 579 wells, and a 450-well synthetic comparison of static-level models with conditional 30-day validation.")
    run.font.size = Pt(10.5)
    run.font.color.rgb = TEXT

    doc.add_page_break()

    h = doc.add_paragraph("Contents", style="Heading 1")
    h.paragraph_format.space_before = Pt(4)
    intro = doc.add_paragraph("The section order is designed for both technical review and machine handoff.")
    intro.runs[0].font.color.rgb = MUTED
    for heading in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.space_after = Pt(2.2)
        if heading.startswith("Appendix"):
            r = p.add_run(heading)
        else:
            number, _, text = heading.partition(". ")
            r = p.add_run(f"{number:>2}  ")
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(BLUE)
            r = p.add_run(text)
        r.font.size = Pt(9.4)
    doc.add_page_break()


def add_equation(doc: Document, equation_lines: list[str]) -> None:
    text = " ".join(line.strip() for line in equation_lines).strip()
    if text.startswith("$$"):
        text = text[2:]
    if text.endswith("$$"):
        text = text[:-2]
    text = text.strip()
    tex = re.sub(r"\s+", " ", text).strip()
    key = hashlib.sha256(tex.encode()).hexdigest()[:16]
    image_dir = Path(os.environ.get("BOREWELL_EQUATION_DIR", "/private/tmp/borewell-equations"))
    manifest = json.loads((image_dir / "manifest.json").read_text())
    p = doc.add_paragraph(style="Equation Display")
    p.add_run().add_picture(str(image_dir / (key + ".png")), width=Inches(manifest[key]["width_inches"]))
    # Equation images retain a searchable source in their accessibility description.
    prop = p.runs[0]._r.find(".//" + qn("wp:docPr"))
    if prop is not None: prop.set("descr", tex)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(r) for r in rows)
    normalized = [r + [""] * (column_count - len(r)) for r in rows]
    table = doc.add_table(rows=len(normalized), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    available_twips = 10000
    widths = [available_twips // column_count] * column_count
    if column_count == 2:
        widths = [3600, 6400]
    elif column_count == 3:
        widths = [4600, 2700, 2700]
    elif column_count == 4:
        widths = [4000, 2000, 2000, 2000]
    elif column_count >= 6:
        widths = [2600] + [(available_twips - 2600) // (column_count - 1)] * (column_count - 1)

    for ridx, data_row in enumerate(normalized):
        row = table.rows[ridx]
        prevent_row_split(row)
        if ridx == 0:
            set_repeat_table_header(row)
        for cidx, value in enumerate(data_row):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_width(cell, widths[cidx] if cidx < len(widths) else widths[-1])
            if ridx == 0:
                set_cell_shading(cell, NAVY)
            elif ridx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, default_bold=(ridx == 0), size=7.8 if column_count >= 6 else 8.3, color=RGBColor(255, 255, 255) if ridx == 0 else TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = parse_table_row(line)
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def is_forced_break_heading(text: str) -> bool:
    return bool(re.match(r"^(18|19|25)\.", text)) or text.startswith("Appendix")


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "Mathematical and Technical Handoff for the Bengaluru Borewell Forecasting Proof of Concept"
    doc.core_properties.subject = "Data, model, validation, rainfall, website, and deployment specification"
    doc.core_properties.author = "Abhiram Srinivas project handoff"
    doc.core_properties.comments = "Generated from the canonical Markdown handoff and visually verified."
    doc.settings.odd_and_even_pages_header_footer = False
    add_cover(doc, headings)

    # Skip title and cover metadata; begin at the first level-two section.
    start = next(i for i, line in enumerate(lines) if line.startswith("## "))
    i = start
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        p = doc.add_paragraph()
        add_inline(p, text)
        paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph(style="Code Block")
                if lines[i].strip():
                    p.add_run(lines[i].rstrip())
                else:
                    p.add_run(" ")
                i += 1
            i += 1
            continue

        if stripped.startswith("$$"):
            flush_paragraph()
            eq = [stripped]
            if not (stripped.endswith("$$") and len(stripped) > 4):
                i += 1
                while i < len(lines):
                    eq.append(lines[i].strip())
                    if lines[i].strip().endswith("$$"):
                        break
                    i += 1
            add_equation(doc, eq)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            flush_paragraph()
            rows = [parse_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            hashes, text = heading_match.groups()
            level = len(hashes) - 1
            forced = ((level == 1 and is_forced_break_heading(text))) and len(doc.paragraphs) > 5
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            p.paragraph_format.page_break_before = forced
            add_inline(p, text, default_bold=True)
            i += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if number_match:
            flush_paragraph()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.72)
            p.paragraph_format.first_line_indent = Cm(-0.42)
            p.paragraph_format.space_after = Pt(2.5)
            prefix = p.add_run(f"{number_match.group(1)}.  ")
            prefix.bold = True
            prefix.font.color.rgb = RGBColor.from_string(BLUE)
            add_inline(p, number_match.group(2))
            i += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet_match.group(1))
            i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()

    # Word compatibility and table border color adjustments.
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.first_child_found_in("w:tblBorders")
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = qn(f"w:{edge}")
            element = borders.find(tag)
            if element is None:
                element = OxmlElement(f"w:{edge}")
                borders.append(element)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), MID_GRAY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
